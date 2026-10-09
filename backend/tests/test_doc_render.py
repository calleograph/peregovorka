"""Единый документ: структура → таблицы Markdown → Word / HTML / текст. Полнота без обрезки, перенос и повтор шапки таблиц, альбомная ориентация для широких таблиц,
мягкие акценты, экранирование, итог и пометка «неполный», структурный режим внешней модели. Без сети."""
from __future__ import annotations

import asyncio
import io
import json

import httpx

from app.services import doc_render as dr
from app.services import export_docs as xd
from app.services import extraction as ex
from app.services.protocols import run_llm_pipeline

from .conftest import login, make_room, make_settings, put_settings, running_app
from .test_admin_features import _drain
from .test_extraction import good, llm_for, reply, TEXT
from .test_transcripts import _feed_and_consume, _join, _segment

LONG = ("Сервер приложений падает при нагрузке выше двухсот запросов в секунду: утечка памяти в обработчике сессий, предложено обновить библиотеку до версии 4.2.1 "
        "и добавить алерт на память; проверить на стенде 192.0.2.55 порт 8443; результат сверить с мониторингом за последние сутки. ") * 3


def item(kind, text, ts, sec, **kw):
    return ex.Item(kind, text, ts, sec, sources=[{"fragment": 1, "ts": ts, "speaker": "Анна Крылова", "quote": text[:40]}], **kw)


def sample_doc(**kw):
    items = [item("topic", "Падение сервера", "10:00:05", 5, detail=LONG), item("decision", "Обновить библиотеку до 4.2.1", "10:05:00", 300),
             item("task", "Проверить на стенде", "10:06:00", 360, assignee="Иван Петров", assignee_id="p02", due="до пятницы", due_date="2026-10-16"),
             item("task", "Настроить алерт", "10:07:00", 420), item("open", "Кто платит за лицензию | и когда", "10:08:00", 480),
             item("topic", "Бюджет на год", "10:20:00", 1200, detail="Согласовали рамки расходов.")]
    return dr.build_document(["Переговорка: Тест", "Дата: 2026-10-12", "Участвовали: Анна Крылова, Иван Петров"], items, summary="Обсудили падение сервера и бюджет.", **kw)


# ------------------------------------------------------------------------------------------------ структура и Markdown
def test_document_has_the_contract_and_tables_keep_full_text_without_ellipsis():
    d = sample_doc()
    assert set(d) >= {"title", "meta", "summary", "discussion", "decisions", "tasks", "proposals", "open_questions", "incomplete", "notes"}
    assert d["discussion"][0] == {"topic": "Падение сервера", "discussion": LONG, "result": "Решение: Обновить библиотеку до 4.2.1. Остался открытым: Кто платит за лицензию | и когда",
                                  "source": "[10:00:05] Анна Крылова"}
    assert d["discussion"][1]["result"] == dr.NO_DECISION, "решения в окне темы нет — так и пишется"
    assert d["tasks"][1]["assignee"] == "не назначен" and d["tasks"][1]["deadline"] == "не указан"
    md = dr.document_to_markdown(d)
    assert LONG.strip() in md and "…" not in md and "..." not in md, "длинная ячейка переносится, а не обрезается"
    assert "| № | Тема | Что обсуждали | Чем закончилось | Где в записи |" in md and "| № | Задача | Ответственный | Срок | Где в записи |" in md
    assert "## Итог" in md and "Обсудили падение сервера и бюджет." in md
    assert "Кто платит за лицензию / и когда" in md, "символ | не ломает таблицу"
    assert "неполным" not in md
    assert "неполным" in dr.document_to_markdown(sample_doc(incomplete=True)) and "[проверить]" in dr.document_to_markdown(sample_doc(incomplete=True))


def test_empty_values_are_honest_not_invented():
    d = dr.build_document(["Переговорка: Тест"], [], summary="")
    md = dr.document_to_markdown(d)
    assert "темы не выделены" in md and "не принималось" in md and "## Итог" not in md


# ------------------------------------------------------------------------------------------------ Word
def test_docx_tables_wrap_repeat_header_and_wide_tables_go_landscape():
    from docx import Document
    from docx.enum.section import WD_ORIENT

    data = xd.to_docx(dr.document_to_markdown(sample_doc()), "Протокол: Тест")
    doc = Document(io.BytesIO(data))
    sec = doc.sections[0]
    assert sec.orientation == WD_ORIENT.LANDSCAPE and sec.page_width > sec.page_height, "таблицы из 5 столбцов — альбомная ориентация"
    usable = sec.page_width - sec.left_margin - sec.right_margin
    assert len(doc.tables) >= 4
    for t in doc.tables:
        xml = t._tbl.xml
        assert "w:tblHeader" in xml and "w:cantSplit" in xml, "шапка повторяется на страницах, строки не рвутся"
        widths = [c.width for c in t.rows[0].cells]
        assert sum(widths) <= usable + 10 and all(w and w > 0 for w in widths), "таблица не выходит за поля"
    assert LONG.strip()[:80] in doc.tables[0].cell(1, 2).text and "…" not in doc.tables[0].cell(1, 2).text
    fills = [t._tbl.xml for t in doc.tables]
    assert any("E3F2E7" in x for x in fills) and any("E3EEFA" in x for x in fills), "решения — зелёный, задачи — голубой акцент"
    # узкие таблицы остаются книжными; код — моноширинным с фоном
    narrow = xd.to_docx("# T\n\n| № | Решение |\n|---|---|\n| 1 | да |\n\n```\nGet-Service | Where-Object Status -eq Running\n```\n")
    d2 = Document(io.BytesIO(narrow))
    assert d2.sections[0].orientation == WD_ORIENT.PORTRAIT
    code_par = next(p for p in d2.paragraphs if "Get-Service" in p.text)
    assert code_par.runs[0].font.name == "Consolas" and "F3F4F6" in code_par._p.xml


# ------------------------------------------------------------------------------------------------ HTML
def test_html_is_self_contained_escaped_and_marks_sections():
    d = sample_doc()
    d["decisions"][0]["decision"] = "<script>alert(1)</script> Обновить"
    page = xd.to_html(dr.document_to_markdown(d), "Протокол <b>x</b>")
    assert page.startswith("<!doctype html>") and "<style>" in page and "<script" not in page.replace("&lt;script", "")
    assert "&lt;script&gt;alert(1)" in page and "<title>Протокол &lt;b&gt;x&lt;/b&gt;</title>" in page
    assert '<table class="ok">' in page and '<table class="info">' in page and "<thead>" in page
    assert LONG.strip() in page.replace("&amp;", "&")
    flag = xd.to_html(dr.document_to_markdown(sample_doc(incomplete=True)))
    assert 'class="flag"' in flag


# ------------------------------------------------------------------------------------------------ структурный режим внешней модели
def test_external_model_in_structured_mode_gets_the_same_document_and_falls_back_when_it_cannot():
    llm, seen = llm_for(good)
    res = asyncio.run(run_llm_pipeline(llm, kind="protocol", instruction="x", text=TEXT, limit=100000, local=None, anonymized=False, external_structured=True))
    assert "## Итог" in res.text and "| № | Задача | Ответственный | Срок | Где в записи |" in res.text and res.structured["document"]["decisions"]
    assert len(seen) == 5 and "response_format" in seen[0]
    # модель не отдаёт разбираемый JSON → обычный режим с предупреждением
    llm2, seen2 = llm_for(lambda b, n: reply("**Протокол**\n\n* пункт") if "response_format" not in b else reply("не JSON"))
    res2 = asyncio.run(run_llm_pipeline(llm2, kind="protocol", instruction="x", text=TEXT, limit=100000, local=None, anonymized=False, external_structured=True))
    assert "**Протокол**" in res2.text and any("Структурный режим не сработал" in w for w in res2.warnings)


def test_summary_failure_leaves_the_document_without_summary_but_complete():
    llm, _ = llm_for(lambda b, n: good(b, n) if "response_format" in b else reply("Оборвано", finish="length"))
    res = asyncio.run(ex.structured_pipeline(llm, kind="protocol", instruction="x", text=TEXT, limit=100000))
    assert "## Итог" not in res.text and "| № | Решение | Где в записи |" in res.text and not res.structured["document"]["incomplete"]


# ------------------------------------------------------------------------------------------------ выгрузка через API
def md_llm(req: httpx.Request) -> httpx.Response:
    return httpx.Response(200, json={"choices": [{"message": {"content": "# Протокол\n\n| № | Решение |\n|---|---|\n| 1 | Утвердили бюджет |"}, "finish_reason": "stop"}]})


def test_protocol_can_be_downloaded_as_html_docx_and_markdown(tmp_path, directory):
    with running_app(make_settings(tmp_path), directory, transports={"llm": httpx.MockTransport(md_llm)}) as c:
        put_settings(c, "llm", enabled=True, provider="external", type="openai_compatible", base_url="http://llm.test/v1", model="m", allow_http=True)
        put_settings(c, "protocol", external_mode="free")
        room = make_room(c)
        a = _join(c, "alice", room["id"])
        _feed_and_consume(c, [_segment(a["meeting_id"], a["identity"], "Утвердили бюджет", offset=5, dur=30)])
        login(c, "alice")
        assert c.post(f"/api/v1/meetings/{a['meeting_id']}/end").status_code == 204
        _drain(c)
        r = c.post(f"/api/v1/meetings/{a['meeting_id']}/protocols", json={"kind": "protocol", "instruction": "x"})
        assert r.status_code == 202, r.text
        _drain(c)
        pid = r.json()["protocol_id"]
        base = f"/api/v1/meetings/{a['meeting_id']}/protocols/{pid}/export"
        h = c.get(base, params={"format": "html"})
        assert h.status_code == 200 and h.headers["content-type"].startswith("text/html") and "Утвердили бюджет" in h.text and "<table" in h.text
        assert c.get(base, params={"format": "docx"}).content[:2] == b"PK" and c.get(base, params={"format": "md"}).status_code == 200
        assert c.get(base, params={"format": "exe"}).status_code == 422


def test_standalone_html_forbids_network_and_scripts_by_its_own_policy():
    page = xd.to_html(dr.document_to_markdown(sample_doc()), "T")
    assert 'http-equiv="Content-Security-Policy"' in page and "default-src 'none'" in page and "script-src" not in page, "документ — без скриптов и сети"
