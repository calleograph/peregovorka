"""Экспорт протоколов: Markdown → простой текст / DOCX / PDF.

Разбор Markdown — собственный, ограниченный (заголовки, абзацы, **жирный**, *курсив*, `код`, списки с вложенностью,
таблицы, цитаты, ссылки, разделители). HTML из ответа модели не интерпретируется: он остаётся обычным текстом.
"""
from __future__ import annotations

import io
import os
import re
from dataclasses import dataclass, field
from pathlib import Path

# ------------------------------------------------------------------------------ разбор
@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False
    href: str | None = None


@dataclass
class ListItem:
    runs: list[Run]
    children: list["Block"] = field(default_factory=list)


@dataclass
class Block:
    kind: str                       # h | p | ul | ol | quote | hr | table | code
    level: int = 0                  # для h
    runs: list[Run] = field(default_factory=list)
    items: list[ListItem] = field(default_factory=list)
    rows: list[list[list[Run]]] = field(default_factory=list)   # table: строки → ячейки → runs
    text: str = ""                  # code


_INLINE = re.compile(r"(\*\*(?P<b>.+?)\*\*|__(?P<b2>.+?)__|\*(?P<i>[^*\s][^*]*?)\*|(?<!\w)_(?P<i2>[^_\s][^_]*?)_(?!\w)|`(?P<c>[^`]+)`|\[(?P<lt>[^\]]+)\]\((?P<lu>[^)\s]+)\))")
_SAFE_URL = re.compile(r"^(https?://|mailto:)", re.I)


def parse_inline(text: str) -> list[Run]:
    runs: list[Run] = []
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            runs.append(Run(text[pos:m.start()]))
        if m.group("b") or m.group("b2"):
            inner = m.group("b") or m.group("b2")
            for r in parse_inline(inner):
                r.bold = True
                runs.append(r)
        elif m.group("i") or m.group("i2"):
            for r in parse_inline(m.group("i") or m.group("i2")):
                r.italic = True
                runs.append(r)
        elif m.group("c"):
            runs.append(Run(m.group("c"), code=True))
        else:
            url = m.group("lu")
            runs.append(Run(m.group("lt"), href=url if _SAFE_URL.match(url) else None))
        pos = m.end()
    if pos < len(text):
        runs.append(Run(text[pos:]))
    return runs


_LI = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")
_HR = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def parse_markdown(md: str) -> list[Block]:
    lines = md.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[Block] = []
    i = 0
    para: list[str] = []

    def flush() -> None:
        nonlocal para
        if para:
            blocks.append(Block("p", runs=parse_inline(" ".join(s.strip() for s in para))))
            para = []

    while i < len(lines):
        line = lines[i]
        if not line.strip():
            flush(); i += 1; continue
        if line.strip().startswith("```"):
            flush(); i += 1; buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            blocks.append(Block("code", text="\n".join(buf)))
            continue
        m = re.match(r"^(#{1,6})\s+(.*?)\s*#*\s*$", line)
        if m:
            flush(); blocks.append(Block("h", level=len(m.group(1)), runs=parse_inline(m.group(2)))); i += 1; continue
        if _HR.match(line):
            flush(); blocks.append(Block("hr")); i += 1; continue
        if line.lstrip().startswith(">"):
            flush(); buf = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                buf.append(lines[i].lstrip()[1:].strip()); i += 1
            blocks.append(Block("quote", runs=parse_inline(" ".join(buf))))
            continue
        if "|" in line and i + 1 < len(lines) and _TABLE_SEP.match(lines[i + 1]):
            flush(); rows = [_split_row(line)]; i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(_split_row(lines[i])); i += 1
            blocks.append(Block("table", rows=[[parse_inline(c) for c in r] for r in rows]))
            continue
        if _LI.match(line):
            flush()
            items: list[tuple[int, bool, str]] = []
            while i < len(lines) and _LI.match(lines[i]):
                mm = _LI.match(lines[i])
                indent, ordered = len(mm.group(1).replace("\t", "    ")), mm.group(2)[0].isdigit()
                if items and indent <= items[0][0] and ordered != items[0][1]:
                    break  # на том же уровне сменился тип списка — это уже другой список
                items.append((indent, ordered, mm.group(3))); i += 1
            blocks.append(_build_list(items))
            continue
        para.append(line); i += 1
    flush()
    return blocks


def _build_list(items: list[tuple[int, bool, str]]) -> Block:
    """Плоский список с отступами → дерево (вложенность по росту отступа)."""
    root = Block("ol" if items[0][1] else "ul")
    stack: list[tuple[int, Block]] = [(items[0][0], root)]
    last: ListItem | None = None
    for indent, ordered, text in items:
        item = ListItem(parse_inline(text))
        if indent > stack[-1][0] and last is not None:
            child = Block("ol" if ordered else "ul")
            last.children.append(child)
            stack.append((indent, child))
        else:
            while len(stack) > 1 and indent < stack[-1][0]:
                stack.pop()
        stack[-1][1].items.append(item)
        last = item
    return root


# ------------------------------------------------------------------------- простой текст
def _runs_text(runs: list[Run]) -> str:
    return "".join(r.text + (f" ({r.href})" if r.href else "") for r in runs)


def md_to_plain(md: str) -> str:
    out: list[str] = []

    def lst(b: Block, depth: int) -> None:
        for n, it in enumerate(b.items, 1):
            marker = f"{n}." if b.kind == "ol" else "•"
            out.append("    " * depth + f"{marker} {_runs_text(it.runs)}")
            for ch in it.children:
                lst(ch, depth + 1)

    for b in parse_markdown(md):
        if b.kind == "h":
            out += [_runs_text(b.runs).upper() if b.level == 1 else _runs_text(b.runs), ""]
        elif b.kind == "p":
            out += [_runs_text(b.runs), ""]
        elif b.kind in ("ul", "ol"):
            lst(b, 0); out.append("")
        elif b.kind == "quote":
            out += ["    " + _runs_text(b.runs), ""]
        elif b.kind == "hr":
            out += ["—" * 20, ""]
        elif b.kind == "table":
            out += ["\t".join(_runs_text(c) for c in r) for r in b.rows] + [""]
        elif b.kind == "code":
            out += [b.text, ""]
    return "\n".join(out).strip() + "\n"


# --------------------------------------------------------------------------------- DOCX
ACCENT_FILL = {"ok": "E3F2E7", "info": "E3EEFA", "warn": "FCEFD0", "muted": "EEF0F3", None: "EEF1F6"}      # мягкие тона: читаются и в чёрно-белой печати
WARN_FILL = "FFF3C4"


def _runs_len(runs: list[Run]) -> int:
    return len("".join(r.text for r in runs))


def column_weights(rows: list[list[list[Run]]], cols: int) -> list[float]:
    """Доли ширины столбцов по содержимому: узкие («№», срок) не растягиваются, длинные (тема, описание) получают место; ничего не обрезается."""
    w = []
    for c in range(cols):
        longest = max((_runs_len(r[c]) for r in rows if c < len(r)), default=4)
        w.append(float(min(max(longest, 5), 70)))
    if cols > 2 and rows and _runs_len(rows[0][0]) <= 2:
        w[0] = min(w[0], 4.0)                                   # столбец «№»
    total = sum(w) or 1.0
    return [x / total for x in w]


def needs_landscape(blocks: list[Block]) -> bool:
    return any(b.kind == "table" and max(len(r) for r in b.rows) >= 5 for b in blocks)


def to_docx(md: str, title: str | None = None) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    from .doc_render import section_accent, warn_text

    blocks = parse_markdown(md)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    sec = doc.sections[0]
    if needs_landscape(blocks):                       # широкие таблицы: альбомная ориентация, а не мелкий нечитаемый текст
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    usable = sec.page_width - sec.left_margin - sec.right_margin
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        try:
            h = doc.styles[name]
            h.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            h.paragraph_format.keep_with_next = True
        except KeyError:
            pass
    if title:
        doc.add_heading(title, level=0)

    def shade(el_pr, fill: str) -> None:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        el_pr.append(shd)

    def add_runs(par, runs: list[Run], size: float | None = None) -> None:
        for r in runs:
            run = par.add_run(r.text + (f" ({r.href})" if r.href else ""))
            run.bold, run.italic = r.bold or None, r.italic or None
            if size:
                run.font.size = Pt(size)
            if r.code:
                run.font.name = "Consolas"

    def add_list(b: Block, depth: int) -> None:
        base = "List Number" if b.kind == "ol" else "List Bullet"
        style = base if depth == 0 else f"{base} {min(depth + 1, 3)}"
        for it in b.items:
            add_runs(doc.add_paragraph(style=style), it.runs)
            for ch in it.children:
                add_list(ch, depth + 1)

    accent: str | None = None
    for b in blocks:
        if b.kind == "h":
            if b.level == 2:
                accent = section_accent(_runs_text(b.runs))
            add_runs(doc.add_heading("", level=min(b.level, 4)), b.runs)
        elif b.kind == "p":
            add_runs(doc.add_paragraph(), b.runs)
        elif b.kind in ("ul", "ol"):
            add_list(b, 0)
        elif b.kind == "quote":
            p = doc.add_paragraph(style="Quote")
            add_runs(p, b.runs)
            if warn_text(_runs_text(b.runs)):
                shade(p._p.get_or_add_pPr(), WARN_FILL)
        elif b.kind == "hr":
            doc.add_paragraph("—" * 30)
        elif b.kind == "code":
            p = doc.add_paragraph()
            r = p.add_run(b.text)
            r.font.name, r.font.size = "Consolas", Pt(9.5)
            p.paragraph_format.keep_together = True
            shade(p._p.get_or_add_pPr(), "F3F4F6")
        elif b.kind == "table":
            cols = max(len(r) for r in b.rows)
            t = doc.add_table(rows=len(b.rows), cols=cols)
            t.style = "Table Grid"
            t.autofit = False
            weights = column_weights(b.rows, cols)
            for ri, row in enumerate(b.rows):
                trpr = t.rows[ri]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))              # строка не разрывается между страницами, если помещается целиком
                if ri == 0:
                    trpr.append(OxmlElement("w:tblHeader"))          # шапка таблицы повторяется на каждой странице
                for ci in range(cols):
                    cell = t.cell(ri, ci)
                    cell.width = int(usable * 0.99 * weights[ci])
                    cell.text = ""
                    par = cell.paragraphs[0]
                    add_runs(par, row[ci] if ci < len(row) else [], size=10)
                    par.paragraph_format.space_after = Pt(2)
                    if ri == 0:
                        for r in par.runs:
                            r.bold = True
                        shade(cell._tc.get_or_add_tcPr(), ACCENT_FILL[accent])
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------------- HTML
HTML_CSS = """
:root{color-scheme:light;--text:#1f2937;--muted:#667085;--line:#d9dee7;--head:#eef1f6;--ok:#e3f2e7;--info:#e3eefa;--warn:#fcefd0;--mute:#eef0f3;--flag:#fff3c4}
*{box-sizing:border-box}body{margin:0;background:#f4f5f8;color:var(--text);font:15px/1.5 system-ui,-apple-system,"Segoe UI",Roboto,Calibri,sans-serif}
main{max-width:1100px;margin:24px auto;padding:28px 34px;background:#fff;border:1px solid var(--line);border-radius:14px}
h1{font-size:26px;margin:0 0 14px;letter-spacing:-.01em}h2{font-size:18px;margin:26px 0 10px;padding-bottom:6px;border-bottom:1px solid var(--line)}h3{font-size:16px}
p{margin:8px 0}em{color:var(--muted)}table{width:100%;border-collapse:collapse;margin:8px 0 14px;table-layout:auto}
th,td{border:1px solid var(--line);padding:7px 10px;text-align:left;vertical-align:top;overflow-wrap:break-word;min-width:9em}th{background:var(--head);font-weight:600}
tr{break-inside:avoid}thead{display:table-header-group}th:first-child,td:first-child{width:2.6em;min-width:0;color:var(--muted)}
.ok th{background:var(--ok)}.info th{background:var(--info)}.warn th{background:var(--warn)}.muted th{background:var(--mute)}
blockquote{margin:10px 0;padding:8px 14px;border-left:4px solid var(--line);background:#fafbfc;color:var(--muted)}blockquote.flag{background:var(--flag);border-color:#e0b100;color:var(--text)}
pre{background:#f3f4f6;border:1px solid var(--line);border-radius:8px;padding:10px 12px;overflow:auto;font:13px/1.45 Consolas,"DejaVu Sans Mono",monospace;white-space:pre-wrap;overflow-wrap:anywhere}
code{font-family:Consolas,"DejaVu Sans Mono",monospace;background:#f3f4f6;padding:1px 4px;border-radius:4px}
@media print{body{background:#fff}main{border:0;margin:0;padding:0;max-width:none}h2{break-after:avoid}}
@page{margin:16mm}
"""


def _esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def _runs_html(runs: list[Run]) -> str:
    out = []
    for r in runs:
        t = _esc(r.text)
        if r.code:
            t = f"<code>{t}</code>"
        if r.bold:
            t = f"<strong>{t}</strong>"
        if r.italic:
            t = f"<em>{t}</em>"
        if r.href:
            t = f'<a href="{_esc(r.href)}" rel="noopener noreferrer">{t}</a>'
        out.append(t)
    return "".join(out)


def to_html(md: str, title: str | None = None) -> str:
    """Самостоятельная страница документа: общие стили, таблицы с переносом, мягкие акценты разделов; печать — через браузер. Текст модели экранируется."""
    from .doc_render import section_accent, warn_text

    def lst(b: Block) -> str:
        tag = "ol" if b.kind == "ol" else "ul"
        return f"<{tag}>" + "".join(f"<li>{_runs_html(it.runs)}" + "".join(lst(c) for c in it.children) + "</li>" for it in b.items) + f"</{tag}>"

    body: list[str] = []
    accent: str | None = None
    for b in parse_markdown(md):
        if b.kind == "h":
            if b.level == 2:
                accent = section_accent(_runs_text(b.runs))
            body.append(f"<h{min(b.level, 4)}>{_runs_html(b.runs)}</h{min(b.level, 4)}>")
        elif b.kind == "p":
            body.append(f"<p>{_runs_html(b.runs)}</p>")
        elif b.kind in ("ul", "ol"):
            body.append(lst(b))
        elif b.kind == "quote":
            cls = ' class="flag"' if warn_text(_runs_text(b.runs)) else ""
            body.append(f"<blockquote{cls}>{_runs_html(b.runs)}</blockquote>")
        elif b.kind == "hr":
            body.append("<hr>")
        elif b.kind == "code":
            body.append(f"<pre>{_esc(b.text)}</pre>")
        elif b.kind == "table":
            head, rows = b.rows[0], b.rows[1:]
            cls = f' class="{accent}"' if accent else ""
            body.append(f"<table{cls}><thead><tr>" + "".join(f"<th>{_runs_html(c)}</th>" for c in head) + "</tr></thead><tbody>"
                        + "".join("<tr>" + "".join(f"<td>{_runs_html(c)}</td>" for c in r) + "</tr>" for r in rows) + "</tbody></table>")
    t = _esc(title or "Документ")
    return (f'<!doctype html>\n<html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">'
            f'<meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'; img-src data:; base-uri \'none\'; form-action \'none\'">'
            f'<title>{t}</title><style>{HTML_CSS}</style></head><body><main>\n' + "\n".join(body) + "\n</main></body></html>\n")


# ---------------------------------------------------------------------------------- PDF
_FONT_DIRS = [os.environ.get("PDF_FONT_DIR", ""), "/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu", "C:/Windows/Fonts",
              "/Library/Fonts", "/System/Library/Fonts/Supplemental"]


def find_fonts() -> dict[str, str] | None:
    """Шрифт с кириллицей: DejaVu (в образе — пакет fonts-dejavu-core) либо Arial как запасной."""
    for family, files in (("DejaVu", ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf", "DejaVuSansMono.ttf")),
                          ("Arial", ("arial.ttf", "arialbd.ttf", "ariali.ttf", "cour.ttf"))):
        for d in _FONT_DIRS:
            if not d:
                continue
            reg = Path(d) / files[0]
            if reg.is_file():
                def pick(name: str, fallback: Path) -> str:
                    p = Path(d) / name
                    return str(p if p.is_file() else fallback)
                return {"family": family, "regular": str(reg), "bold": pick(files[1], reg), "italic": pick(files[2], reg), "mono": pick(files[3], reg)}
    return None


def to_pdf(md: str, title: str | None = None) -> bytes:
    from fpdf import FPDF

    fonts = find_fonts()
    if fonts is None:
        raise RuntimeError("Не найден шрифт с кириллицей для PDF (установите fonts-dejavu-core или задайте PDF_FONT_DIR)")
    pdf = FPDF(format="A4")
    pdf.set_margins(18, 16, 18)
    pdf.set_auto_page_break(True, 16)
    fam = "Body"
    pdf.add_font(fam, "", fonts["regular"])
    pdf.add_font(fam, "B", fonts["bold"])
    pdf.add_font(fam, "I", fonts["italic"])
    pdf.add_font("Mono", "", fonts["mono"])
    pdf.add_page()

    def write_runs(runs: list[Run], size: float = 11, h: float = 6) -> None:
        for r in runs:
            style = ("B" if r.bold else "") + ("I" if r.italic else "")
            if r.code:
                pdf.set_font("Mono", "", size - 1)
            else:
                pdf.set_font(fam, style, size)
            pdf.write(h, r.text + (f" ({r.href})" if r.href else ""))
        pdf.set_font(fam, "", size)

    def newline(h: float = 6) -> None:
        pdf.ln(h)

    if title:
        pdf.set_font(fam, "B", 16)
        pdf.multi_cell(0, 8, title, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)

    def pdf_list(b: Block, depth: int) -> None:
        for n, it in enumerate(b.items, 1):
            pdf.set_x(pdf.l_margin + 5 * depth)
            marker = f"{n}. " if b.kind == "ol" else "•  "
            pdf.set_font(fam, "", 11)
            pdf.write(6, marker)
            write_runs(it.runs)
            newline(6)
            for ch in it.children:
                pdf_list(ch, depth + 1)

    for b in parse_markdown(md):
        if b.kind == "h":
            size = {1: 15, 2: 13.5, 3: 12.5}.get(b.level, 11.5)
            pdf.ln(2)
            write_runs([Run(r.text, bold=True, italic=r.italic, href=r.href) for r in b.runs], size, size * 0.55)
            newline(size * 0.6); pdf.ln(1)
        elif b.kind == "p":
            write_runs(b.runs); newline(6.5)
        elif b.kind in ("ul", "ol"):
            pdf_list(b, 0); pdf.ln(1)
        elif b.kind == "quote":
            pdf.set_x(pdf.l_margin + 6)
            write_runs([Run(r.text, bold=r.bold, italic=True, href=r.href) for r in b.runs]); newline(6.5)
        elif b.kind == "hr":
            y = pdf.get_y() + 2
            pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y); pdf.ln(5)
        elif b.kind == "code":
            pdf.set_font("Mono", "", 9.5)
            pdf.multi_cell(0, 5, b.text, new_x="LMARGIN", new_y="NEXT"); pdf.set_font(fam, "", 11); pdf.ln(2)
        elif b.kind == "table":
            pdf.set_font(fam, "", 10)
            with pdf.table(first_row_as_headings=True, text_align="LEFT") as table:
                for row in b.rows:
                    r = table.row()
                    for cell in row:
                        r.cell(_runs_text(cell))
            pdf.set_font(fam, "", 11); pdf.ln(3)
    return bytes(pdf.output())
